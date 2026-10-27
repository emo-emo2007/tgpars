import argparse
import asyncio
import os
import re
import sys
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import requests
from requests.adapters import HTTPAdapter
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


TELEGRAM_HOSTS = {"t.me", "www.t.me", "telegram.me", "www.telegram.me"}


@dataclass(frozen=True)
class TelegramPost:
    url: str
    text: str
    photo_urls: tuple[str, ...]
    photo_data: tuple[bytes, ...] = ()


def parse_post_url(url: str) -> tuple[str, str]:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or parsed.hostname not in TELEGRAM_HOSTS:
        raise ValueError("Ожидается публичная ссылка вида https://t.me/channel/123")

    parts = [part for part in parsed.path.split("/") if part]
    if parts and parts[0] == "s":
        parts = parts[1:]
    if parts and parts[0] == "c":
        if len(parts) == 3 and parts[1].isdigit() and parts[2].isdigit():
            return f"-100{parts[1]}", parts[2]
        raise ValueError("Не удалось определить приватный канал и номер поста из ссылки")
    if len(parts) != 2 or not parts[1].isdigit():
        raise ValueError("Не удалось определить публичный канал и номер поста из ссылки")
    return parts[0], parts[1]


def _message_text(message) -> str:
    text_node = message.select_one(".tgme_widget_message_text")
    if text_node is None:
        return ""
    for line_break in text_node.find_all("br"):
        line_break.replace_with("\n")
    lines = [line.strip() for line in text_node.get_text().replace("\xa0", " ").splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _message_photos(message) -> tuple[str, ...]:
    photo_urls = []
    for photo in message.select(".tgme_widget_message_photo_wrap"):
        style = photo.get("style", "")
        match = re.search(r"background-image\s*:\s*url\(['\"]?(.*?)['\"]?\)", style, re.IGNORECASE)
        if match:
            photo_url = match.group(1).strip()
            if photo_url.startswith("//"):
                photo_url = f"https:{photo_url}"
            if photo_url and photo_url not in photo_urls:
                photo_urls.append(photo_url)
    return tuple(photo_urls)


def extract_post(html: str, username: str, message_id: str, source_url: str) -> TelegramPost:
    soup = BeautifulSoup(html, "html.parser")
    expected_id = message_id
    for message in soup.select(".tgme_widget_message[data-post]"):
        if message.get("data-post", "").rsplit("/", 1)[-1] == expected_id:
            return TelegramPost(source_url, _message_text(message), _message_photos(message))
    raise LookupError(
        f"Пост {username}/{message_id} не найден в публичной ленте. "
        "Возможно, канал приватный, пост удалён или Telegram не показывает его без авторизации."
    )


def fetch_post(session: requests.Session, url: str) -> TelegramPost:
    username, message_id = parse_post_url(url)
    if username.startswith("-100"):
        raise ValueError("Ссылка на приватный канал требует запуска с --telegram-auth")
    page_url = f"https://t.me/s/{username}?before={int(message_id) + 1}"
    response = session.get(page_url, timeout=30)
    response.raise_for_status()
    return extract_post(response.text, username, message_id, url)


def create_http_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 (compatible; TelegramPostExporter/1.0)"})
    session.mount("https://", HTTPAdapter(max_retries=3))
    session.mount("http://", HTTPAdapter(max_retries=3))
    return session


async def fetch_posts_with_login(urls: list[str]) -> tuple[list[TelegramPost], int]:
    api_id = os.getenv("TELEGRAM_API_ID")
    api_hash = os.getenv("TELEGRAM_API_HASH")
    if not api_id or not api_hash:
        raise RuntimeError(
            "Для входа задайте переменные TELEGRAM_API_ID и TELEGRAM_API_HASH. "
            "Их можно получить на https://my.telegram.org."
        )
    try:
        api_id_number = int(api_id)
    except ValueError as error:
        raise RuntimeError("TELEGRAM_API_ID должен быть числом.") from error

    try:
        from telethon import TelegramClient
    except ImportError as error:
        raise RuntimeError("Установите зависимости командой: python -m pip install -r requirements.txt") from error

    session_path = Path.home() / ".telegram_post_exporter"
    posts = []
    failures = 0
    async with TelegramClient(str(session_path), api_id_number, api_hash) as client:
        await client.start()
        for url in urls:
            try:
                username, message_id = parse_post_url(url)
                entity_reference = int(username) if username.startswith("-100") else username
                entity = await client.get_entity(entity_reference)
                message = await client.get_messages(entity, ids=int(message_id))
                if message is None:
                    raise LookupError("Пост не найден или недоступен этому аккаунту.")
                photo_data = ()
                if message.photo:
                    image = await client.download_media(message, file=bytes)
                    if image:
                        photo_data = (image,)
                posts.append(TelegramPost(url, message.message or "", (), photo_data))
                print(f"Получен пост: {url}")
            except Exception as error:
                failures += 1
                print(f"Не удалось обработать {url}: {error}", file=sys.stderr)
    return posts, failures


def _set_run_font(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def create_document(posts: list[TelegramPost], output_path: Path, session: requests.Session) -> None:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")

    for index, post in enumerate(posts):
        if index:
            document.add_page_break()

        link_paragraph = document.add_paragraph()
        _set_run_font(link_paragraph.add_run(post.url))

        for photo_url in post.photo_urls:
            try:
                response = session.get(photo_url, timeout=(10, 30))
                response.raise_for_status()
                photo_paragraph = document.add_paragraph()
                photo_paragraph.add_run().add_picture(BytesIO(response.content), width=Inches(6.1))
            except (requests.RequestException, ValueError) as error:
                print(f"Не удалось добавить фото {photo_url}: {error}", file=sys.stderr)
        for image in post.photo_data:
            photo_paragraph = document.add_paragraph()
            photo_paragraph.add_run().add_picture(BytesIO(image), width=Inches(6.1))

        text_paragraph = document.add_paragraph()
        _set_run_font(text_paragraph.add_run(post.text))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def read_links(input_path: Path) -> list[str]:
    return [
        line.strip()
        for line in input_path.read_text(encoding="utf-8-sig").splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Собрать публичные посты Telegram в документ Word")
    parser.add_argument("input", type=Path, help="Текстовый файл: одна ссылка на пост в каждой строке")
    parser.add_argument("-o", "--output", type=Path, default=Path("telegram_posts.docx"), help="Путь к DOCX")
    parser.add_argument(
        "--telegram-auth",
        action="store_true",
        help="Получать посты через авторизованный Telegram-аккаунт",
    )
    args = parser.parse_args()

    try:
        links = read_links(args.input)
    except OSError as error:
        print(f"Не удалось прочитать файл ссылок: {error}", file=sys.stderr)
        return 1
    if not links:
        print("В файле нет ссылок на посты.", file=sys.stderr)
        return 1

    session = create_http_session()
    if args.telegram_auth:
        try:
            posts, failures = asyncio.run(fetch_posts_with_login(links))
        except RuntimeError as error:
            print(error, file=sys.stderr)
            return 1
    else:
        posts = []
        failures = 0
        for url in links:
            try:
                posts.append(fetch_post(session, url))
                print(f"Получен пост: {url}")
            except (ValueError, LookupError, requests.RequestException) as error:
                failures += 1
                print(f"Не удалось обработать {url}: {error}", file=sys.stderr)

    if not posts:
        print("Документ не создан: ни один пост не удалось получить.", file=sys.stderr)
        return 1

    try:
        create_document(posts, args.output, session)
    except OSError as error:
        print(f"Не удалось сохранить документ: {error}", file=sys.stderr)
        return 1
    print(f"Сохранено постов: {len(posts)}. Файл: {args.output}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
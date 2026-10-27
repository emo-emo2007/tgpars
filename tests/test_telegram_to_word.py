import unittest
import tempfile
from pathlib import Path

import requests
from docx import Document

from telegram_to_word import TelegramPost, create_document, extract_post, parse_post_url


class ParsePostUrlTests(unittest.TestCase):
    def test_accepts_regular_and_preview_links(self):
        self.assertEqual(parse_post_url("https://t.me/tr_sputnik/175574"), ("tr_sputnik", "175574"))
        self.assertEqual(parse_post_url("https://t.me/s/tr_sputnik/175574?single"), ("tr_sputnik", "175574"))

    def test_parses_private_channel_links_for_authorized_mode(self):
        self.assertEqual(parse_post_url("https://t.me/c/12345/678"), ("-10012345", "678"))


class ExtractPostTests(unittest.TestCase):
    def test_extracts_text_and_photo_from_public_post_markup(self):
        html = """
        <div class="tgme_widget_message" data-post="sample_channel/42">
          <div class="tgme_widget_message_text" dir="auto">
            🌍 Putin: Afrika geleceğin kıtasıdır<br/><br/><b>Rusya</b>, Afrika hakkında.
          </div>
          <a class="tgme_widget_message_photo_wrap" style="background-image:url('https://cdn.example/photo.jpg')"></a>
        </div>
        """

        post = extract_post(html, "sample_channel", "42", "https://t.me/sample_channel/42")

        self.assertEqual(post.text, "🌍 Putin: Afrika geleceğin kıtasıdır\n\nRusya, Afrika hakkında.")
        self.assertEqual(post.photo_urls, ("https://cdn.example/photo.jpg",))

    def test_reports_when_target_post_is_not_in_page(self):
        with self.assertRaisesRegex(LookupError, "не найден"):
            extract_post("<html></html>", "sample_channel", "42", "https://t.me/sample_channel/42")


class WordDocumentTests(unittest.TestCase):
    def test_writes_link_and_text_with_requested_font(self):
        post = TelegramPost("https://t.me/sample_channel/42", "Текст статьи", ())
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "posts.docx"
            create_document([post], output_path, requests.Session())
            document = Document(output_path)

        self.assertEqual(document.paragraphs[0].text, post.url)
        self.assertEqual(document.paragraphs[1].text, post.text)
        self.assertEqual(document.paragraphs[1].runs[0].font.name, "Times New Roman")
        self.assertEqual(document.paragraphs[1].runs[0].font.size.pt, 12)


if __name__ == "__main__":
    unittest.main()